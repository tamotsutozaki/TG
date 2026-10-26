# -*- coding: utf-8 -*-
"""
Gera a monografia (TG) formatada em ABNT, a partir do texto em monografia/conteudo/*.md
e das figuras em monografia/figuras/, seguindo o modelo de tópicos (modelo.pdf) e as
normas da Fatec Indaiatuba (Manual-de-Normas-2021.pdf).

- Formatação: A4; margens 3/2/3/2 cm; Times New Roman 12; espaço 1,5; recuo de 1a linha
  1,25 cm; justificado; sem espaço entre parágrafos; títulos 14/12 negrito, separados do
  texto por uma linha de 1,5.
- Ilustrações: legenda acima ("Figura 1 – ..."), fonte abaixo, tamanho 10. Quadros
  (dados qualitativos) com todas as bordas; tabelas (dados numéricos) sem bordas laterais.
- Listas de figuras, quadros e tabelas e o sumário são campos do Word, atualizados ao
  final pelo próprio Word (atualizar_campos.ps1), que também exporta o PDF.

Marcação aceita nos arquivos .md:
  # TÍTULO              título sem numeração, centralizado, em nova página
  ## 1 SEÇÃO            seção primária     ### 1.1 Seção / #### 1.1.1 Seção  subseções
  a) item               alínea              - item          item de lista
  *itálico*  **negrito**  [[PREENCHER: ...]] (destacado em amarelo)
  @fig:id @qua:id @tab:id  referência cruzada ("Figura 3")
  {{figura id=".." arquivo=".." largura="15" titulo=".." fonte=".."}}
  {{quadro id=".." titulo=".." fonte=".." larguras="3,5,8" nota=".."}} | a | b | ... {{/quadro}}
  {{tabela ...}} ... {{/tabela}}
"""
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(AQUI, "TG_Pedro_Tozaki.docx")
PDF = os.path.join(AQUI, "TG_Pedro_Tozaki.pdf")
DIR_CONT = os.path.join(AQUI, "monografia", "conteudo")
DIR_FIG = os.path.join(AQUI, "monografia", "figuras")

AUTOR = "PEDRO HENRIQUE TAMOTSU TOZAKI"
TITULO = ("Sistema web para gestão do fluxo de trabalho em laboratórios de "
          "patologia veterinária: da solicitação à emissão de laudos")
CIDADE = "INDAIATUBA"
ANO = "2026"
ORIENTACAO = "Marlon [[PREENCHER: titulação e sobrenome]]"
LARGURA_TEXTO = 16.0                 # cm (21 - 3 - 2)
ESPACO_TITULO = Pt(18)               # uma linha em branco de 1,5 (12 pt x 1,5)
CAPITULOS = ["01-introducao.md", "02-capitulo-1.md", "03-capitulo-2.md", "04-capitulo-3.md",
             "05-conclusao.md", "06-referencias.md", "07-apendices.md"]
ROTULOS = {"fig": "Figura", "qua": "Quadro", "tab": "Tabela"}


# ---------------------------------------------------------------------------
# Estilos
# ---------------------------------------------------------------------------
def fonte_times(elemento_rpr):
    rf = elemento_rpr.get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(attr), "Times New Roman")


def set_base_style(document):
    st = document.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    fonte_times(st.element.get_or_add_rPr())
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    # idioma pt-BR (corretor ortográfico) no padrão do documento e no Normal
    rpr_default = document.styles.element.find(qn("w:docDefaults")).find(
        qn("w:rPrDefault")).find(qn("w:rPr"))
    for alvo in (rpr_default, st.element.get_or_add_rPr()):
        lang = alvo.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            alvo.append(lang)
        lang.set(qn("w:val"), "pt-BR")
        lang.set(qn("w:eastAsia"), "pt-BR")

    # títulos: Times, preto (o modelo padrão vem azul/Calibri)
    for nome, tam in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
        h = document.styles[nome]
        h.font.name = "Times New Roman"
        h.font.size = Pt(tam)
        h.font.bold = True
        h.font.italic = False
        h.font.color.rgb = RGBColor(0, 0, 0)
        fonte_times(h.element.get_or_add_rPr())
        h.paragraph_format.space_before = Pt(0)
        h.paragraph_format.space_after = ESPACO_TITULO
        h.paragraph_format.keep_with_next = True


# ---------------------------------------------------------------------------
# Texto com marcação simples
# ---------------------------------------------------------------------------
TOKENS = re.compile(r"(\*\*.+?\*\*|\*.+?\*|\[\[.+?\]\])")


def _run(p, text, bold=False, italic=False, size=12, destaque=False):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    fonte_times(r._element.get_or_add_rPr())
    if destaque:
        hl = OxmlElement("w:highlight")
        hl.set(qn("w:val"), "yellow")
        r._element.get_or_add_rPr().append(hl)
    return r


def escrever(p, texto, size=12, bold=False, italic=False):
    """Escreve o texto no parágrafo interpretando *itálico*, **negrito** e [[destaque]]."""
    for parte in TOKENS.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
            escrever(p, parte[2:-2], size, True, italic)
        elif parte.startswith("[[") and parte.endswith("]]"):
            _run(p, "[" + parte[2:-2] + "]", bold=True, italic=italic, size=size, destaque=True)
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            escrever(p, parte[1:-1], size, bold, True)
        else:
            _run(p, parte, bold, italic, size)


def formatar(p, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, recuo=1.25, esquerda=0.0, simples=False,
             antes=0, depois=0, manter=False):
    pf = p.paragraph_format
    p.alignment = alinhamento
    pf.first_line_indent = Cm(recuo)
    pf.left_indent = Cm(esquerda)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE if simples else WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(antes)
    pf.space_after = Pt(depois)
    pf.keep_with_next = manter
    return p


def add_field(paragraph, instr, resultado="", size=None):
    r = paragraph.add_run()
    if size:
        r.font.size = Pt(size)
    fb = OxmlElement("w:fldChar")
    fb.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    fs = OxmlElement("w:fldChar")
    fs.set(qn("w:fldCharType"), "separate")
    tx = OxmlElement("w:t")
    tx.text = resultado
    fe = OxmlElement("w:fldChar")
    fe.set(qn("w:fldCharType"), "end")
    for el in (fb, it, fs, tx, fe):
        r._r.append(el)
    return r


# ---------------------------------------------------------------------------
# Leitura dos arquivos de conteúdo
# ---------------------------------------------------------------------------
def ler(nome):
    with open(os.path.join(DIR_CONT, nome), encoding="utf-8") as f:
        return f.read()


def atributos(linha):
    return dict(re.findall(r'(\w+)="([^"]*)"', linha))


def blocos(texto):
    linhas = texto.splitlines()
    i, out = 0, []
    especial = re.compile(r"^(#{1,4} |\{\{|[a-z]\) |- |> )")
    while i < len(linhas):
        l = linhas[i].rstrip()
        if not l.strip():
            i += 1
            continue
        if l.startswith("#### "):
            out.append(("h3", l[5:].strip()))
        elif l.startswith("### "):
            out.append(("h2", l[4:].strip()))
        elif l.startswith("## "):
            out.append(("h1", l[3:].strip()))
        elif l.startswith("# "):
            out.append(("titulo", l[2:].strip()))
        elif l.startswith("{{figura"):
            out.append(("figura", atributos(l)))
        elif l.startswith("{{quadro") or l.startswith("{{tabela"):
            tipo = "quadro" if l.startswith("{{quadro") else "tabela"
            attrs, linhas_tab = atributos(l), []
            i += 1
            while not linhas[i].strip().startswith("{{/" + tipo):
                if linhas[i].strip().startswith("|"):
                    celulas = [c.strip() for c in linhas[i].strip().strip("|").split("|")]
                    linhas_tab.append(celulas)
                i += 1
            out.append((tipo, attrs, linhas_tab))
        elif l.strip() == "{{quebra}}":
            out.append(("quebra", None))
        elif re.match(r"^[a-z]\) ", l):
            itens = []
            while i < len(linhas) and re.match(r"^[a-z]\) ", linhas[i]):
                itens.append(linhas[i].strip())
                i += 1
            out.append(("alineas", itens))
            continue
        elif l.startswith("- "):
            itens = []
            while i < len(linhas) and linhas[i].startswith("- "):
                itens.append(linhas[i][2:].strip())
                i += 1
            out.append(("itens", itens))
            continue
        elif l.startswith("> "):
            out.append(("citacao", l[2:].strip()))
        else:
            partes = []
            while i < len(linhas) and linhas[i].strip() and not especial.match(linhas[i]):
                partes.append(linhas[i].strip())
                i += 1
            out.append(("par", " ".join(partes)))
            continue
        i += 1
    return out


def numerar(todos_blocos):
    """Primeira passagem: número de cada figura/quadro/tabela, na ordem em que aparecem."""
    cont = {"figura": 0, "quadro": 0, "tabela": 0}
    numeros = {}
    for b in todos_blocos:
        if b[0] in cont:
            cont[b[0]] += 1
            chave = {"figura": "fig", "quadro": "qua", "tabela": "tab"}[b[0]]
            numeros[f"{chave}:{b[1]['id']}"] = cont[b[0]]
    return numeros


def referencias_cruzadas(texto, numeros):
    def troca(m):
        chave = f"{m.group(1)}:{m.group(2)}"
        if chave not in numeros:
            raise KeyError(f"Referência inexistente: @{chave}")
        return f"{ROTULOS[m.group(1)]} {numeros[chave]}"
    return re.sub(r"@(fig|qua|tab):([\w-]+)", troca, texto)


# ---------------------------------------------------------------------------
# Elementos do documento
# ---------------------------------------------------------------------------
def titulo_sem_numero(doc, texto, nova_pagina=True, toc=True):
    p = doc.add_paragraph(style="Heading 1") if toc else doc.add_paragraph()
    formatar(p, WD_ALIGN_PARAGRAPH.CENTER, recuo=0, depois=18)
    p.paragraph_format.page_break_before = nova_pagina
    p.paragraph_format.keep_with_next = True
    escrever(p, texto, size=14, bold=True)
    return p


def titulo_secao(doc, texto, nivel):
    estilo = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[nivel]
    p = doc.add_paragraph(style=estilo)
    formatar(p, WD_ALIGN_PARAGRAPH.LEFT, recuo=0, antes=18, depois=18, manter=True)
    escrever(p, texto, size=14 if nivel == 1 else 12, bold=True)
    return p


def legenda(doc, rotulo, numero, titulo, manter=True):
    p = doc.add_paragraph()
    formatar(p, WD_ALIGN_PARAGRAPH.CENTER, recuo=0, simples=True, antes=12, depois=4, manter=manter)
    _run(p, f"{rotulo} ", size=10)
    add_field(p, f"SEQ {rotulo} \\* ARABIC", str(numero), size=10)
    _run(p, " – ", size=10)
    escrever(p, titulo, size=10)
    return p


def fonte_nota(doc, fonte, nota=None, alinhamento=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    formatar(p, alinhamento, recuo=0, simples=True, antes=4, depois=0 if nota else 12)
    escrever(p, "Fonte: " + fonte if not fonte.startswith("Fonte") else fonte, size=10)
    if nota:
        p2 = doc.add_paragraph()
        formatar(p2, alinhamento, recuo=0, simples=True, antes=2, depois=12)
        escrever(p2, "Nota: " + nota, size=10)


def figura(doc, attrs, numero):
    legenda(doc, "Figura", numero, attrs["titulo"])
    p = doc.add_paragraph()
    formatar(p, WD_ALIGN_PARAGRAPH.CENTER, recuo=0, simples=True, manter=True)
    caminho = os.path.join(DIR_FIG, attrs["arquivo"])
    p.add_run().add_picture(caminho, width=Cm(float(attrs.get("largura", LARGURA_TEXTO))))
    fonte_nota(doc, attrs.get("fonte", "Elaborado pelo autor (2026)."))


def borda(el_pai, lados, val="single", sz="4", cor="000000"):
    for lado in lados:
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), val)
        b.set(qn("w:sz"), sz)
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), cor)
        el_pai.append(b)


def sombrear(celula, cor):
    tcpr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    tcpr.append(shd)


def tabela(doc, tipo, attrs, linhas, numero, numeros):
    rotulo = "Quadro" if tipo == "quadro" else "Tabela"
    legenda(doc, rotulo, numero, attrs["titulo"])
    ncol = len(linhas[0])
    larguras = [float(x) for x in attrs.get("larguras", ",".join(["1"] * ncol)).split(",")]
    escala = LARGURA_TEXTO / sum(larguras)
    larguras = [w * escala for w in larguras]

    t = doc.add_table(rows=len(linhas), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False  # layout fixo: respeita as larguras das colunas
    tblpr = t._tbl.tblPr
    # o esquema do Word exige a ordem dos filhos de tblPr
    bordas = OxmlElement("w:tblBorders")
    if tipo == "quadro":
        borda(bordas, ("top", "left", "bottom", "right", "insideH", "insideV"))
    else:
        borda(bordas, ("top", "bottom"), sz="8")
    tblpr.insert_element_before(bordas, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                                "w:tblCaption", "w:tblDescription")
    mar = OxmlElement("w:tblCellMar")  # margens internas das células
    for lado, valor in (("top", "20"), ("left", "80"), ("bottom", "20"), ("right", "80")):
        m = OxmlElement(f"w:{lado}")
        m.set(qn("w:w"), valor)
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tblpr.insert_element_before(mar, "w:tblLook", "w:tblCaption", "w:tblDescription")

    for r, celulas in enumerate(linhas):
        row = t.rows[r]
        trpr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:cantSplit")
        trpr.append(el)
        if r == 0:
            el = OxmlElement("w:tblHeader")
            trpr.append(el)
        for c in range(ncol):
            cel = row.cells[c]
            cel.width = Cm(larguras[c])
            if r == 0:
                if tipo == "quadro":
                    sombrear(cel, "F2F2F2")
                else:
                    tcpr = cel._tc.get_or_add_tcPr()
                    tb = OxmlElement("w:tcBorders")
                    borda(tb, ("bottom",))
                    tcpr.append(tb)
            p = cel.paragraphs[0]
            centro = tipo == "tabela" and c > 0
            # todas as linhas "presas" à seguinte: a tabela não se separa da fonte
            formatar(p, WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.LEFT,
                     recuo=0, simples=True, manter=True)
            texto = referencias_cruzadas(celulas[c] if c < len(celulas) else "", numeros)
            escrever(p, texto, size=10, bold=(r == 0))
    fonte_nota(doc, attrs.get("fonte", "Elaborado pelo autor (2026)."), attrs.get("nota"),
               alinhamento=WD_ALIGN_PARAGRAPH.LEFT)


def lista_itens(doc, itens, numeros, marcador=None):
    for item in itens:
        p = doc.add_paragraph()
        recuo_pendente = 0.6
        formatar(p, WD_ALIGN_PARAGRAPH.JUSTIFY, recuo=-recuo_pendente, esquerda=1.25 + recuo_pendente)
        escrever(p, (marcador + " " if marcador else "") + referencias_cruzadas(item, numeros))


def corpo(doc, conteudo, numeros):
    primeiro = True
    em_referencias = False
    for b in conteudo:
        tipo = b[0]
        if tipo == "titulo":
            titulo_sem_numero(doc, b[1], nova_pagina=not primeiro)
            primeiro = False
            em_referencias = b[1].upper().startswith("REFERÊNCIAS")
        elif tipo == "h1":
            titulo_secao(doc, b[1], 1)
        elif tipo == "h2":
            titulo_secao(doc, b[1], 2)
        elif tipo == "h3":
            titulo_secao(doc, b[1], 3)
        elif tipo == "par":
            p = doc.add_paragraph()
            if em_referencias:
                formatar(p, WD_ALIGN_PARAGRAPH.LEFT, recuo=0, simples=True, depois=12)
            else:
                formatar(p)
            escrever(p, referencias_cruzadas(b[1], numeros))
        elif tipo == "alineas":
            lista_itens(doc, b[1], numeros)
        elif tipo == "itens":
            lista_itens(doc, b[1], numeros, marcador="–")
        elif tipo == "citacao":
            p = doc.add_paragraph()
            formatar(p, recuo=0, esquerda=4, simples=True, antes=12, depois=12)
            escrever(p, b[1], size=10)
        elif tipo == "figura":
            figura(doc, b[1], numeros[f"fig:{b[1]['id']}"])
        elif tipo in ("quadro", "tabela"):
            chave = "qua" if tipo == "quadro" else "tab"
            tabela(doc, tipo, b[1], b[2], numeros[f"{chave}:{b[1]['id']}"], numeros)
        elif tipo == "quebra":
            doc.add_page_break()


# ---------------------------------------------------------------------------
# Montagem do documento
# ---------------------------------------------------------------------------
def margens(sec):
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(3)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(3)
    sec.right_margin = Cm(2)


def linha_central(doc, texto, bold=True):
    p = doc.add_paragraph()
    formatar(p, WD_ALIGN_PARAGRAPH.CENTER, recuo=0)
    escrever(p, texto, bold=bold)


def vazio(doc, n):
    for _ in range(n):
        formatar(doc.add_paragraph(), recuo=0)


def main():
    doc = Document()
    set_base_style(doc)
    margens(doc.sections[0])

    cab = ["CENTRO ESTADUAL DE EDUCAÇÃO TECNOLÓGICA PAULA SOUZA",
           "FACULDADE DE TECNOLOGIA DE INDAIATUBA",
           "DR. ARCHIMEDES LAMOGLIA",
           "CURSO DE TECNOLOGIA",
           "EM ANÁLISE E DESENVOLVIMENTO DE SISTEMAS"]

    # ---- CAPA ----
    for l in cab:
        linha_central(doc, l)
    vazio(doc, 6)
    linha_central(doc, AUTOR)
    vazio(doc, 2)
    linha_central(doc, TITULO)
    vazio(doc, 10)
    for l in (CIDADE, ANO):
        linha_central(doc, l)

    # ---- FOLHA DE ROSTO: nova seção; a contagem de páginas começa aqui (a capa não conta) ----
    doc.add_section(WD_SECTION.NEW_PAGE)
    margens(doc.sections[-1])
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    doc.sections[-1]._sectPr.insert_element_before(
        pg, "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg", "w:textDirection",
        "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange")
    for l in cab:
        linha_central(doc, l)
    vazio(doc, 4)
    linha_central(doc, AUTOR)
    vazio(doc, 2)
    linha_central(doc, TITULO)
    vazio(doc, 3)
    p = doc.add_paragraph()
    formatar(p, WD_ALIGN_PARAGRAPH.JUSTIFY, recuo=0, esquerda=8, simples=True)
    escrever(p, "Trabalho de Graduação apresentado por Pedro Henrique Tamotsu Tozaki como "
                "pré-requisito parcial para a conclusão do Curso Superior de Tecnologia em Análise e "
                "Desenvolvimento de Sistemas, da Faculdade de Tecnologia de Indaiatuba, elaborado sob a "
                f"orientação do(a) Prof.(a) {ORIENTACAO}.")
    vazio(doc, 8)
    for l in (CIDADE, ANO):
        linha_central(doc, l)
    doc.add_page_break()

    # ---- RESUMO ----
    resumo = [x.strip() for x in ler("00-resumo.md").split("\n\n") if x.strip()]
    titulo_sem_numero(doc, "RESUMO", nova_pagina=False, toc=False)
    for par in resumo:
        p = doc.add_paragraph()
        formatar(p, WD_ALIGN_PARAGRAPH.JUSTIFY, recuo=0, simples=True, depois=12)
        if par.startswith("Palavras-chave:"):
            _run(p, "Palavras-chave: ", bold=True)
            escrever(p, par.split(":", 1)[1].strip())
        else:
            escrever(p, par)
    doc.add_page_break()

    # ---- LISTAS ----
    for titulo, rotulo in (("LISTA DE FIGURAS", "Figura"), ("LISTA DE QUADROS", "Quadro"),
                           ("LISTA DE TABELAS", "Tabela")):
        titulo_sem_numero(doc, titulo, nova_pagina=False, toc=False)
        p = doc.add_paragraph()
        formatar(p, WD_ALIGN_PARAGRAPH.LEFT, recuo=0)
        add_field(p, f'TOC \\h \\z \\c "{rotulo}"')
        doc.add_page_break()

    titulo_sem_numero(doc, "LISTA DE SIGLAS", nova_pagina=False, toc=False)
    for l in ler("00-siglas.md").splitlines():
        if l.strip():
            p = doc.add_paragraph()
            formatar(p, WD_ALIGN_PARAGRAPH.LEFT, recuo=0)
            escrever(p, l.strip())
    doc.add_page_break()

    # ---- SUMÁRIO ----
    titulo_sem_numero(doc, "SUMÁRIO", nova_pagina=False, toc=False)
    p = doc.add_paragraph()
    formatar(p, WD_ALIGN_PARAGRAPH.LEFT, recuo=0)
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u')

    # ---- PARTE TEXTUAL: nova seção, paginação visível a partir daqui ----
    doc.add_section(WD_SECTION.NEW_PAGE)
    tsec = doc.sections[-1]
    margens(tsec)
    # a nova seção copia a anterior: remove o reinício da numeração para continuar a contagem
    for pg in tsec._sectPr.findall(qn("w:pgNumType")):
        tsec._sectPr.remove(pg)
    tsec.header.is_linked_to_previous = False
    hp = tsec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.first_line_indent = Cm(0)
    add_field(hp, "PAGE", "1")
    # capa e pré-textuais sem número visível
    for sec in doc.sections[:-1]:
        sec.header.is_linked_to_previous = False

    conteudo = []
    for nome in CAPITULOS:
        conteudo += blocos(ler(nome))
    numeros = numerar(conteudo)
    corpo(doc, conteudo, numeros)

    doc.save(OUT)
    print("Gerado:", OUT)

    # Word atualiza sumário, listas e numeração das legendas e exporta o PDF
    if "--sem-word" not in sys.argv and sys.platform == "win32":
        script = os.path.join(AQUI, "monografia", "atualizar_campos.ps1")
        r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", script,
                            "-Docx", OUT, "-Pdf", PDF], capture_output=True, text=True)
        print(r.stdout.strip() or r.stderr.strip())


if __name__ == "__main__":
    main()
